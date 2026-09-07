import os
import io
import re
import copy
import html
import calendar
from datetime import date
from typing import Dict, List, Any, Optional
import openpyxl
from openpyxl.cell.rich_text import TextBlock, CellRichText
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Default files in workspace
import sys

def get_base_dir():
    if getattr(sys, 'frozen', False):
        return getattr(sys, '_MEIPASS', os.path.dirname(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))

WORKSPACE_DIR = get_base_dir()
DEFAULT_MATRIX_PATH = os.path.join(WORKSPACE_DIR, "2026_MA TRAN GIAM SAT TAI FECT_2026_thaotdp_update 04.9.2026.xlsx")
DEFAULT_TEMPLATE_PATH = os.path.join(WORKSPACE_DIR, "template_Ke hoach kiem soat, giam sat_team QAFECT.xlsx")

# Standard thin border
THIN_SIDE = Side(border_style="thin", color="D3D3D3")
MEDIUM_SIDE = Side(border_style="thin", color="000000")
CELL_BORDER = Border(left=MEDIUM_SIDE, right=MEDIUM_SIDE, top=MEDIUM_SIDE, bottom=MEDIUM_SIDE)

DATA_FONT = Font(name="Arial", size=10, bold=False)
TITLE_FONT = Font(name="Arial", size=14, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True)


def cell_to_html(cell) -> str:
    """
    Converts cell content (including CellRichText, bold, italic, font colors) to HTML
    for rich display in the web preview table.
    """
    val = cell.value
    if val is None:
        return ""

    if isinstance(val, CellRichText):
        html_chunks = []
        for part in val:
            if isinstance(part, TextBlock):
                txt = html.escape(part.text)
                styles = []
                if part.font:
                    if part.font.bold:
                        styles.append("font-weight:bold")
                    if part.font.italic:
                        styles.append("font-style:italic")
                    if part.font.color and part.font.color.rgb:
                        rgb_str = str(part.font.color.rgb)
                        if len(rgb_str) == 8:
                            rgb_str = rgb_str[2:]
                        styles.append(f"color:#{rgb_str}")
                style_str = f' style="{"; ".join(styles)}"' if styles else ""
                html_chunks.append(f"<span{style_str}>{txt}</span>")
            else:
                html_chunks.append(html.escape(str(part)))
        return "".join(html_chunks)

    escaped = html.escape(str(val))
    styles = []
    if cell.font:
        if cell.font.bold:
            styles.append("font-weight:bold")
        if cell.font.italic:
            styles.append("font-style:italic")
        if cell.font.color and cell.font.color.rgb:
            rgb_str = str(cell.font.color.rgb)
            if len(rgb_str) == 8:
                rgb_str = rgb_str[2:]
            styles.append(f"color:#{rgb_str}")
    if styles:
        return f'<span style="{"; ".join(styles)}">{escaped}</span>'
    return escaped


def is_unit_sheet(ws) -> Optional[Dict[str, int]]:
    """
    Check if a sheet is a unit matrix sheet by inspecting row 2 for month headers (T1..T12).
    Returns a mapping of {'T1': col_idx, 'T2': col_idx, ...} if valid, else None.
    """
    # Check row 2 (or row 1)
    month_cols = {}
    for r in [2, 1]:
        for col_idx in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=col_idx).value
            if val is not None:
                val_str = str(val).strip().upper()
                m = re.match(r"^T(1[0-2]|[1-9])$", val_str)
                if m:
                    month_cols[val_str] = col_idx
        if len(month_cols) >= 6:  # Found at least 6 months
            return month_cols
    return None


def get_matrix_info(matrix_source=None) -> Dict[str, Any]:
    """
    Parses matrix file and returns sheet units and task counts for all 12 months.
    """
    if matrix_source is None:
        matrix_source = DEFAULT_MATRIX_PATH
    elif isinstance(matrix_source, bytes):
        matrix_source = io.BytesIO(matrix_source)

    wb = openpyxl.load_workbook(matrix_source, data_only=True, read_only=True)
    units_info = {}
    months_list = [f"T{i}" for i in range(1, 13)]
    month_totals = {m: 0 for m in months_list}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        month_cols = is_unit_sheet(ws)
        if not month_cols:
            continue

        # Count tasks per month
        unit_counts = {m: 0 for m in months_list}
        total_unit_tasks = 0

        # Scan rows starting from row 3
        for row in ws.iter_rows(min_row=3, values_only=True):
            # Must have at least document or content or detail
            doc = str(row[1] or "").strip() if len(row) > 1 else ""
            content = str(row[2] or "").strip() if len(row) > 2 else ""
            detail = str(row[3] or "").strip() if len(row) > 3 else ""
            if not (doc or content or detail):
                continue

            # Check if row is a total/summary row
            col_a = str(row[0] or "").strip().lower()
            if "tổng" in col_a or "cộng" in col_a or "tổng" in doc.lower():
                continue

            for m_key in months_list:
                c_idx = month_cols.get(m_key)
                if c_idx and c_idx <= len(row):
                    cell_val = row[c_idx - 1]
                    if cell_val is not None:
                        val_str = str(cell_val).strip().lower()
                        if val_str == "x":
                            unit_counts[m_key] += 1
                            month_totals[m_key] += 1
                            total_unit_tasks += 1

        units_info[sheet_name] = {
            "name": sheet_name,
            "counts": unit_counts,
            "total": total_unit_tasks,
        }

    wb.close()
    return {
        "units": units_info,
        "month_totals": month_totals,
        "all_months": months_list,
    }


def extract_tasks_for_month(matrix_source=None, target_month: str = "T8") -> Dict[str, List[Dict[str, Any]]]:
    """
    Extracts all tasks for the specified month (e.g. 'T8' or '8') from all unit sheets.
    """
    if matrix_source is None:
        matrix_source = DEFAULT_MATRIX_PATH
    elif isinstance(matrix_source, bytes):
        matrix_source = io.BytesIO(matrix_source)

    m_normalized = target_month.strip().upper()
    if not m_normalized.startswith("T"):
        m_normalized = f"T{m_normalized}"

    wb = openpyxl.load_workbook(matrix_source, rich_text=True)
    results = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        month_cols = is_unit_sheet(ws)
        if not month_cols:
            continue

        target_col = month_cols.get(m_normalized)
        if not target_col:
            continue

        sheet_tasks = []
        for r_idx in range(3, ws.max_row + 1):
            col_a = str(ws.cell(row=r_idx, column=1).value or "").strip()
            cell_doc = ws.cell(row=r_idx, column=2)
            cell_content = ws.cell(row=r_idx, column=3)
            cell_detail = ws.cell(row=r_idx, column=4)
            cell_dept = ws.cell(row=r_idx, column=5)

            doc_str = str(cell_doc.value or "").strip()
            content_str = str(cell_content.value or "").strip()
            detail_str = str(cell_detail.value or "").strip()
            dept_str = str(cell_dept.value or "").strip()

            if not (doc_str or content_str or detail_str):
                continue

            if "tổng" in col_a.lower() or "tổng" in doc_str.lower() or "cộng" in col_a.lower():
                continue

            # Check mark in target month column
            month_val = ws.cell(row=r_idx, column=target_col).value
            if month_val is not None and str(month_val).strip().lower() == "x":
                sheet_tasks.append({
                    "row_index": r_idx,
                    "unit": sheet_name,
                    "department": dept_str,
                    "document": doc_str,
                    "content": content_str,
                    "detail": detail_str,
                    "department_html": cell_to_html(cell_dept),
                    "document_html": cell_to_html(cell_doc),
                    "content_html": cell_to_html(cell_content),
                    "detail_html": cell_to_html(cell_detail),
                    # Raw objects for exact preservation in exported Excel
                    "doc_val": copy.deepcopy(cell_doc.value),
                    "doc_font": copy.copy(cell_doc.font) if cell_doc.has_style and cell_doc.font else None,
                    "doc_fill": copy.copy(cell_doc.fill) if cell_doc.has_style and cell_doc.fill and cell_doc.fill.patternType and cell_doc.fill.patternType != 'none' else None,
                    "content_val": copy.deepcopy(cell_content.value),
                    "content_font": copy.copy(cell_content.font) if cell_content.has_style and cell_content.font else None,
                    "content_fill": copy.copy(cell_content.fill) if cell_content.has_style and cell_content.fill and cell_content.fill.patternType and cell_content.fill.patternType != 'none' else None,
                    "detail_val": copy.deepcopy(cell_detail.value),
                    "detail_font": copy.copy(cell_detail.font) if cell_detail.has_style and cell_detail.font else None,
                    "detail_fill": copy.copy(cell_detail.fill) if cell_detail.has_style and cell_detail.fill and cell_detail.fill.patternType and cell_detail.fill.patternType != 'none' else None,
                    "dept_val": copy.deepcopy(cell_dept.value),
                    "dept_font": copy.copy(cell_dept.font) if cell_dept.has_style and cell_dept.font else None,
                    "dept_fill": copy.copy(cell_dept.fill) if cell_dept.has_style and cell_dept.fill and cell_dept.fill.patternType and cell_dept.fill.patternType != 'none' else None,
                })

        results[sheet_name] = sheet_tasks

    wb.close()
    return results


def update_calendar_sheet(ws, month_num: int, year: int = 2026):
    """
    Updates the calendar grid sheet (e.g. LICH KS_T8.2026) for the chosen month and year.
    Fills in the dates for Mondays to Fridays.
    """
    # Title in A2
    month_2digit = f"{month_num:02d}"
    ws.cell(row=2, column=1).value = f"Tháng {month_2digit}.{year}"

    # Find working days (Mon-Fri) grouped by weeks, filtering out weeks with no Mon-Fri days
    cal = calendar.Calendar(firstweekday=calendar.MONDAY)
    month_days = cal.monthdayscalendar(year, month_num)
    working_weeks = [week for week in month_days if any(d > 0 for d in week[:5])]

    # Date rows in template: Row 4 (W1), Row 7 (W2), Row 10 (W3), Row 13 (W4), Row 16 (W5)
    # Columns B (T2) -> F (T6) correspond to weekday indexes 0 -> 4
    date_rows = [4, 7, 10, 13, 16]
    for w_idx, week in enumerate(working_weeks[:5]):
        r_num = date_rows[w_idx] if w_idx < len(date_rows) else 16
        for d_idx in range(5):  # Mon=0, Tue=1, Wed=2, Thu=3, Fri=4
            col_letter = get_column_letter(2 + d_idx)  # Col B is Mon
            day_val = week[d_idx]
            if day_val > 0:
                d_str = f"{day_val:02d}.{month_num}.{year}"
                ws[f"{col_letter}{r_num}"] = d_str
            else:
                ws[f"{col_letter}{r_num}"] = ""


def generate_control_plan_workbook(
    matrix_source=None,
    template_source=None,
    month_num: int = 8,
    year: int = 2026,
    fsc_mode: str = "fsc_ct",  # "fsc_ct" or "all_fsc" (creates FSC CT, FSC ST, FSC HG)
    include_calendar: bool = True,
    selected_units: Optional[List[str]] = None,
    include_empty_sheets: bool = True,
    auto_schedule: bool = True,
) -> io.BytesIO:
    """
    Generates the new Excel workbook populated from the matrix file according to template.
    Returns BytesIO object containing the xlsx file.
    """
    if matrix_source is None:
        matrix_source = DEFAULT_MATRIX_PATH
    elif isinstance(matrix_source, bytes):
        matrix_source = io.BytesIO(matrix_source)

    if template_source is None:
        template_source = DEFAULT_TEMPLATE_PATH
    elif isinstance(template_source, bytes):
        template_source = io.BytesIO(template_source)

    month_str = f"T{month_num}"
    month_2digit = f"{month_num:02d}"

    # Extract tasks
    tasks_by_unit = extract_tasks_for_month(matrix_source, month_str)

    # Load template workbook
    wb_template = openpyxl.load_workbook(template_source, rich_text=True)

    # Map unit sheets:
    # Matrix unit sheet -> Target template sheet name(s)
    # Matrix has: 'VP FE', 'FPTU', 'FGW', 'FSB', 'FSW', 'FSC'
    # Template has: 'VP FE', 'FPTU', 'FGW', 'FSW', 'FSC CT'
    from scheduler import get_auditor, schedule_tasks, populate_calendar_grid

    # Determine sheets to produce
    target_units_map = {}
    for matrix_unit, task_list in tasks_by_unit.items():
        if selected_units and matrix_unit not in selected_units:
            continue

        if matrix_unit == "FSC":
            if fsc_mode == "all_fsc":
                for fsc_sub in ["FSC CT", "FSC ST", "FSC HG"]:
                    sub_tasks = []
                    for t in task_list:
                        t_copy = dict(t)
                        t_copy["unit"] = fsc_sub
                        t_copy["scheduled_auditor"] = get_auditor(fsc_sub, t.get("department", ""))
                        sub_tasks.append(t_copy)
                    target_units_map[fsc_sub] = ("FSC", sub_tasks)
            else:
                ct_tasks = []
                for t in task_list:
                    t_copy = dict(t)
                    t_copy["unit"] = "FSC CT"
                    t_copy["scheduled_auditor"] = get_auditor("FSC CT", t.get("department", ""))
                    ct_tasks.append(t_copy)
                target_units_map["FSC CT"] = ("FSC", ct_tasks)
        else:
            u_tasks = []
            for t in task_list:
                t_copy = dict(t)
                t_copy["unit"] = matrix_unit
                t_copy["scheduled_auditor"] = get_auditor(matrix_unit, t.get("department", ""))
                u_tasks.append(t_copy)
            target_units_map[matrix_unit] = (matrix_unit, u_tasks)

    # Auto-schedule tasks if requested
    calendar_events = []
    if auto_schedule:
        sched_input = {k: v[1] for k, v in target_units_map.items()}
        sched_output, calendar_events = schedule_tasks(sched_input, year, month_num)
        for k in target_units_map.keys():
            orig_src = target_units_map[k][0]
            target_units_map[k] = (orig_src, sched_output.get(k, []))

    # Update calendar sheet
    cal_sheet = None
    for s_name in wb_template.sheetnames:
        if "LICH KS" in s_name.upper():
            cal_sheet = wb_template[s_name]
            cal_sheet.title = f"LICH KS_T{month_num}.{year}"
            update_calendar_sheet(cal_sheet, month_num, year)
            if auto_schedule:
                populate_calendar_grid(cal_sheet, calendar_events, year, month_num)
            break

    if not include_calendar and cal_sheet:
        wb_template.remove(cal_sheet)

    # Unit sheets processing
    # Collect sample style sheet from template to copy for units not yet in template (e.g. FSB)
    sample_sheet = None
    for s_name in ["FPTU", "FGW", "FSW", "VP FE", "FSC CT"]:
        if s_name in wb_template.sheetnames:
            sample_sheet = wb_template[s_name]
            break

    # Process each target unit
    for target_name, (matrix_unit, tasks) in target_units_map.items():
        if not include_empty_sheets and len(tasks) == 0:
            if target_name in wb_template.sheetnames:
                wb_template.remove(wb_template[target_name])
            continue

        # Get or create worksheet
        if target_name in wb_template.sheetnames:
            ws = wb_template[target_name]
        else:
            # Create sheet by copying sample sheet
            if sample_sheet is not None:
                ws = wb_template.copy_worksheet(sample_sheet)
                ws.title = target_name
            else:
                ws = wb_template.create_sheet(title=target_name)

        # Update Title in A1
        ws["A1"] = f"KẾ HOẠCH KIỂM SOÁT THÁNG {month_2digit}. {year}"
        ws["A1"].font = TITLE_FONT

        # Template prototype row styles from row 4
        proto_styles = {}
        for c_idx in range(1, 11):
            proto_cell = ws.cell(row=4, column=c_idx)
            proto_styles[c_idx] = {
                "font": Font(name=proto_cell.font.name or "Arial", size=proto_cell.font.size or 10, bold=proto_cell.font.bold),
                "alignment": Alignment(
                    horizontal="center" if c_idx in (1, 8, 9, 10) else "left",
                    vertical="top",
                    wrap_text=True
                ),
                "border": CELL_BORDER,
                "fill": PatternFill(fill_type=None),
            }

        # Clear existing data rows (from row 4 down)
        if ws.max_row >= 4:
            ws.delete_rows(4, ws.max_row - 3)

        # Populate rows preserving rich text, font colors, bold, italic, and fills from matrix
        for t_idx, task in enumerate(tasks):
            r_num = 4 + t_idx

            # Col A: Đơn vị
            cell_a = ws.cell(row=r_num, column=1, value=target_name)
            cell_a.font = DATA_FONT
            cell_a.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            cell_a.border = CELL_BORDER

            # Col B: Phòng/ CB (from Dept in matrix)
            cell_b = ws.cell(row=r_num, column=2, value=copy.deepcopy(task["dept_val"]))
            cell_b.font = copy.copy(task["dept_font"]) if task["dept_font"] else DATA_FONT
            if task["dept_fill"]:
                cell_b.fill = copy.copy(task["dept_fill"])
            cell_b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell_b.border = CELL_BORDER

            # Col C: Đại diện đơn vị được đánh giá (blank)
            cell_c = ws.cell(row=r_num, column=3, value="")
            cell_c.font = DATA_FONT
            cell_c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell_c.border = CELL_BORDER

            # Col D: Tài liệu (from Document in matrix)
            cell_d = ws.cell(row=r_num, column=4, value=copy.deepcopy(task["doc_val"]))
            cell_d.font = copy.copy(task["doc_font"]) if task["doc_font"] else DATA_FONT
            if task["doc_fill"]:
                cell_d.fill = copy.copy(task["doc_fill"])
            cell_d.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell_d.border = CELL_BORDER

            # Col E: Nội dung yêu cầu (from Content in matrix)
            cell_e = ws.cell(row=r_num, column=5, value=copy.deepcopy(task["content_val"]))
            cell_e.font = copy.copy(task["content_font"]) if task["content_font"] else DATA_FONT
            if task["content_fill"]:
                cell_e.fill = copy.copy(task["content_fill"])
            cell_e.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell_e.border = CELL_BORDER

            # Col F: Chi tiết yêu cầu (from Detail in matrix)
            cell_f = ws.cell(row=r_num, column=6, value=copy.deepcopy(task["detail_val"]))
            cell_f.font = copy.copy(task["detail_font"]) if task["detail_font"] else DATA_FONT
            if task["detail_fill"]:
                cell_f.fill = copy.copy(task["detail_fill"])
            cell_f.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell_f.border = CELL_BORDER

            # Col G: Thời gian phát sinh hồ sơ (blank)
            cell_g = ws.cell(row=r_num, column=7, value="")
            cell_g.font = DATA_FONT
            cell_g.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            cell_g.border = CELL_BORDER

            # Col H: Ngày
            cell_h = ws.cell(row=r_num, column=8, value=task.get("scheduled_date", "") if auto_schedule else "")
            cell_h.font = DATA_FONT
            cell_h.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            cell_h.border = CELL_BORDER

            # Col I: Giờ
            cell_i = ws.cell(row=r_num, column=9, value=task.get("scheduled_time", "") if auto_schedule else "")
            cell_i.font = DATA_FONT
            cell_i.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            cell_i.border = CELL_BORDER

            # Col J: CB đánh giá
            auditor_val = task.get("scheduled_auditor", "") or get_auditor(target_sheet_name, task.get("department", ""))
            cell_j = ws.cell(row=r_num, column=10, value=auditor_val)
            cell_j.font = DATA_FONT
            cell_j.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            cell_j.border = CELL_BORDER

    # Remove any template unit sheets that were not in target_units_map
    for s_name in list(wb_template.sheetnames):
        if "LICH KS" in s_name.upper():
            continue
        if s_name not in target_units_map:
            wb_template.remove(wb_template[s_name])

    # Save to BytesIO
    output = io.BytesIO()
    wb_template.save(output)
    output.seek(0)
    return output
