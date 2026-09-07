import re
import io
import datetime
from collections import defaultdict
from typing import Dict, List, Any, Optional
import openpyxl

def extract_representatives(raw_text: Any) -> List[str]:
    """
    Extracts individual personnel names or account codes from Column C text.
    Handles delimiters: semicolons, commas, newlines, +, &, and ' và '/' and '.
    """
    if not raw_text:
        return []
    text = str(raw_text).replace("\r", "\n")
    parts = re.split(r"[;,\n+&]|(?:\s+và\s+)|(?:\s+and\s+)", text, flags=re.IGNORECASE)
    cleaned = []
    for p in parts:
        p = p.strip()
        p = re.sub(r"^[\s\-\*•]+", "", p).strip()
        if p and p not in cleaned:
            cleaned.append(p)
    return cleaned


def rep_to_email(rep_str: str) -> str:
    """
    Converts a representative account string to an email address (@fe.edu.vn).
    Supports:
    - Raw email: e.g. kimttt3@fe.edu.vn
    - Account code: e.g. Thainvt -> thainvt@fe.edu.vn
    - Parenthesized account: e.g. Nguyễn Thị Lan (LanNT12) -> lannt12@fe.edu.vn
    - Prefixed account: e.g. HauNT51 - P. Đào Tạo -> haunt51@fe.edu.vn
    """
    if not rep_str:
        return ""
    s = rep_str.strip()
    if "@" in s:
        return s.lower()
    m_paren = re.search(r"\(([A-Za-z0-9_.\-]+)\)", s)
    if m_paren:
        acc = m_paren.group(1).strip()
        return f"{acc.lower()}@fe.edu.vn"
    m_prefix = re.match(r"^([A-Za-z0-9_.\-]+)\s*[\-:]", s)
    if m_prefix:
        acc = m_prefix.group(1).strip()
        return f"{acc.lower()}@fe.edu.vn"
    if " " not in s:
        clean_acc = re.sub(r"[^\w.]", "", s)
        if clean_acc:
            return f"{clean_acc.lower()}@fe.edu.vn"
    tokens = s.split()
    for tok in reversed(tokens):
        clean = re.sub(r"[^\w.]", "", tok)
        if re.match(r"^[A-Za-z]+[0-9]*$", clean) and len(clean) >= 3:
            return f"{clean.lower()}@fe.edu.vn"
    return ""


def parse_date_list(val: Any) -> List[datetime.date]:
    """
    Parses date representations from Excel into a list of datetime.date.
    Supports single date, multiple dates (e.g. '24 & 25.8.2026', '21, 22.8.2026'),
    and date ranges (e.g. '14.8.2026 -> 18.8.2026').
    """
    if isinstance(val, datetime.datetime):
        return [val.date()]
    if isinstance(val, datetime.date):
        return [val]
    if not val:
        return []
    s = str(val).strip()
    s = re.sub(r"\s+", " ", s)

    # 1. Range: dd.mm.yyyy -> dd.mm.yyyy
    m_range = re.search(r"(\d{1,2})[./\-](\d{1,2})[./\-](\d{4})\s*(?:->|-|đến)\s*(\d{1,2})[./\-](\d{1,2})[./\-](\d{4})", s)
    if m_range:
        d1, m1, y1 = int(m_range.group(1)), int(m_range.group(2)), int(m_range.group(3))
        d2, m2, y2 = int(m_range.group(4)), int(m_range.group(5)), int(m_range.group(6))
        try:
            start_date = datetime.date(y1, m1, d1)
            if y2 > y1 + 1:  # Protect against typo years like 2028
                y2 = y1
            end_date = datetime.date(y2, m2, d2)
            if start_date <= end_date and (end_date - start_date).days <= 31:
                dates = []
                cur = start_date
                while cur <= end_date:
                    dates.append(cur)
                    cur += datetime.timedelta(days=1)
                return dates
        except ValueError:
            pass

    # 2. Multiple days: d1 & d2.m.y or d1, d2, d3.m.y
    m_multi = re.match(r"^([\d\s&,/+và]+)[./\-](\d{1,2})[./\-](\d{4})$", s)
    if m_multi:
        days_part = m_multi.group(1)
        month = int(m_multi.group(2))
        year = int(m_multi.group(3))
        day_strs = re.findall(r"\b\d{1,2}\b", days_part)
        results = []
        for ds in day_strs:
            try:
                results.append(datetime.date(year, month, int(ds)))
            except ValueError:
                pass
        if results:
            return results

    # 3. Standard single date: dd.mm.yyyy or yyyy-mm-dd
    m_single = re.search(r"(\d{1,2})[./\-](\d{1,2})[./\-](\d{4})", s)
    if m_single:
        try:
            return [datetime.date(int(m_single.group(3)), int(m_single.group(2)), int(m_single.group(1)))]
        except ValueError:
            pass

    m_iso = re.search(r"(\d{4})[./\-](\d{1,2})[./\-](\d{1,2})", s)
    if m_iso:
        try:
            return [datetime.date(int(m_iso.group(1)), int(m_iso.group(2)), int(m_iso.group(3)))]
        except ValueError:
            pass

    return []


def parse_plan_excel(file_source: Any) -> Dict[str, Any]:
    """
    Parses a finalized QA FECT Plan Excel file.
    Supports both file path and in-memory BytesIO/bytes.
    Extracts scheduled events from unit sheets (col 8, 9, 10, col 3 representative) and calendar grid.
    """
    if isinstance(file_source, bytes):
        file_source = io.BytesIO(file_source)

    wb = openpyxl.load_workbook(file_source, data_only=True)

    # 1. Identify Calendar Sheet
    cal_sheet_name = None
    month_detected = ""
    year_detected = 2026
    for s_name in wb.sheetnames:
        if "LICH" in s_name.upper():
            cal_sheet_name = s_name
            m = re.search(r"T(\d{1,2})", s_name)
            if m:
                month_detected = f"T{m.group(1)}"
            y = re.search(r"20\d{2}", s_name)
            if y:
                year_detected = int(y.group(0))
            break

    # 2. Extract detailed tasks from Unit Sheets
    unit_sheets = []
    tasks_by_date = defaultdict(lambda: defaultdict(list))
    total_task_count = 0

    for s_name in wb.sheetnames:
        if "LICH" in s_name.upper():
            continue
        ws = wb[s_name]
        unit_sheets.append(s_name)

        # Detect the specific header row
        header_row = None
        for r in range(1, min(ws.max_row + 1, 10)):
            row_text = " ".join([str(ws.cell(row=r, column=c).value or "").lower() for c in range(1, 15)])
            if "đơn vị" in row_text and ("ngày" in row_text or "cb đánh giá" in row_text):
                header_row = r
                break

        if not header_row:
            header_row = 3

        cols = {
            "unit": 1,
            "dept": 2,
            "representative": 3,
            "doc": 4,
            "content": 5,
            "detail": 6,
            "record_time": 7,
            "date": 8,
            "time": 9,
            "auditor": 10,
        }

        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(row=header_row, column=c).value or "").strip().lower()
            if val == "ngày":
                cols["date"] = c
            elif val in ["giờ", "thời gian", "thời gian đánh giá"]:
                cols["time"] = c
            elif "cb đánh giá" in val or "người đánh giá" in val or "cb dg" in val:
                cols["auditor"] = c
            elif "đại diện" in val or "được đánh giá" in val or "nhân sự" in val:
                cols["representative"] = c
            elif "phát sinh" in val or ("hồ sơ" in val and "thời gian" in val) or val in ["thời gian phát sinh hồ sơ", "thời gian phát sinh"]:
                cols["record_time"] = c
            elif val in ["đơn vị", "đơn vị được đánh giá"]:
                cols["unit"] = c
            elif "phòng" in val:
                cols["dept"] = c
            elif "tài liệu" in val:
                cols["doc"] = c
            elif "nội dung" in val:
                cols["content"] = c
            elif "chi tiết" in val:
                cols["detail"] = c

        # Read tasks row by row
        start_r = header_row + 1
        for r in range(start_r, ws.max_row + 1):
            d_val = ws.cell(row=r, column=cols["date"]).value
            aud_val = ws.cell(row=r, column=cols["auditor"]).value
            if not d_val or not aud_val:
                continue

            date_objs = parse_date_list(d_val)
            if not date_objs:
                continue

            aud_str = re.sub(r"ThanhNTD(?!6)", "ThanhNTD6", str(aud_val).strip())
            rep_raw = ws.cell(row=r, column=cols["representative"]).value
            rep_str = str(rep_raw or "").strip()
            reps_list = extract_representatives(rep_str)

            unit_val = str(ws.cell(row=r, column=cols["unit"]).value or s_name).strip()
            dept_val = str(ws.cell(row=r, column=cols["dept"]).value or "").strip()
            doc_val = str(ws.cell(row=r, column=cols["doc"]).value or "").strip()
            content_val = str(ws.cell(row=r, column=cols["content"]).value or "").strip()
            detail_val = str(ws.cell(row=r, column=cols["detail"]).value or "").strip()
            time_val = str(ws.cell(row=r, column=cols["time"]).value or "").strip()

            record_time_col = cols.get("record_time", 7)
            record_time_raw = ws.cell(row=r, column=record_time_col).value
            record_time_val = str(record_time_raw or "").strip()
            if record_time_val.lower() == "none":
                record_time_val = ""

            for d_obj in date_objs:
                d_str = f"{d_obj.day:02d}.{d_obj.month:02d}.{d_obj.year}"
                tasks_by_date[d_str][aud_str].append({
                    "unit": unit_val,
                    "department": dept_val,
                    "representative_raw": rep_str,
                    "representatives": reps_list,
                    "document": doc_val,
                    "content": content_val,
                    "detail": detail_val,
                    "record_time": record_time_val,
                    "time": time_val,
                    "auditor": aud_str,
                    "date_obj": d_obj,
                    "date_str": d_str,
                })
                total_task_count += 1

    # 3. Group into Scheduled Events by (Date, Auditor)
    events_list = []
    events_by_date = {}

    sorted_dates = sorted(
        tasks_by_date.keys(),
        key=lambda ds: (int(ds.split(".")[2]), int(ds.split(".")[1]), int(ds.split(".")[0]))
    )

    for d_str in sorted_dates:
        aud_map = tasks_by_date[d_str]
        date_events = []
        for aud, items in aud_map.items():
            units = list(dict.fromkeys(it["unit"] for it in items if it["unit"]))
            depts = list(dict.fromkeys(it["department"] for it in items if it["department"]))
            docs = list(dict.fromkeys(it["document"] for it in items if it["document"]))
            contents = list(dict.fromkeys(it["content"] for it in items if it["content"]))
            time_str = items[0]["time"] if items else ""
            d_obj = items[0]["date_obj"] if items else None

            # Collect representatives and convert to emails (Cột C: Đại diện đơn vị được đánh giá)
            all_reps = []
            for it in items:
                for rep in it.get("representatives", []):
                    if rep and rep not in all_reps:
                        all_reps.append(rep)

            rep_emails = []
            for rep in all_reps:
                em = rep_to_email(rep)
                if em and em not in rep_emails:
                    rep_emails.append(em)

            # Collect record times (Cột G: Thời gian phát sinh hồ sơ)
            record_times = []
            for it in items:
                rt = it.get("record_time", "")
                if rt and rt not in record_times:
                    record_times.append(rt)

            record_time_display = "; ".join(record_times) if record_times else ""

            # Format representatives display with @fe.edu.vn accounts (Cột C)
            formatted_reps = [rep_to_email(r) or r for r in all_reps]
            reps_display = ", ".join(formatted_reps) if formatted_reps else "Chưa có thông tin"

            # Determine session
            session_str = "Cả ngày"
            if "12h00" in time_str:
                session_str = "Sáng"
            elif "13h00" in time_str or "13h30" in time_str:
                session_str = "Chiều"

            ev = {
                "id": f"{d_str}_{aud}_{session_str}",
                "date": d_str,
                "date_iso": d_obj.strftime("%Y-%m-%d") if d_obj else "",
                "date_obj": d_obj,
                "auditor": aud,
                "time": time_str or ("09h00 - 17h00" if session_str == "Cả ngày" else ("09h00 - 12h00" if session_str == "Sáng" else "13h00 - 17h30")),
                "session": session_str,
                "units": units,
                "unit_display": " & ".join(units),
                "departments": depts,
                "dept_display": "; ".join(depts),
                "representatives": formatted_reps,
                "representatives_display": reps_display,
                "representative_emails": rep_emails,
                "record_times": record_times,
                "record_time_display": record_time_display,
                "documents": docs,
                "contents": contents,
                "task_count": len(items),
                "tasks": items,
            }
            date_events.append(ev)
            events_list.append(ev)

        events_by_date[d_str] = date_events

    return {
        "success": True,
        "month": month_detected or (f"T{sorted_dates[0].split('.')[1]}" if sorted_dates else "T8"),
        "year": year_detected,
        "calendar_sheet": cal_sheet_name,
        "unit_sheets": unit_sheets,
        "total_tasks": total_task_count,
        "total_events": len(events_list),
        "dates": sorted_dates,
        "events": events_list,
        "events_by_date": events_by_date,
    }
